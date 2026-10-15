import json, shutil
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT=Path('..').resolve(); EV=Path('evidence').resolve()
SOURCE=ROOT/'DGCP_NORTIC_B2_A_Evaluacion_Marcada.xlsx'
OUTPUT=ROOT/'DGCP_NORTIC_B2_A_Auditoria_Tecnica_Final_2026-10-14.xlsx'
REPORT=ROOT/'Informe_Auditoria_Tecnica_DGCP_NORTIC_B2_A_2026-10-14.md'

S={
'1.04.':('NO CUMPLE','La política POL.DGCP.DC.002 existe, pero no satisface íntegramente 1.04.c y falta evidencia contractual para 1.04.f. Evidencia: documento adjunto, págs. 3, 5-8; E-10.'),
'1.04.a.':('CUMPLE','POL.DGCP.DC.002 es un documento independiente, aprobado y firmado, denominado “Gestión y Accesibilidad de portales Web”. Evidencia: págs. 1-2 y 8; E-10.'),
'1.04.b.':('CUMPLE','POL.DGCP.DC.002 §3.9 asigna al webmaster la supervisión y garantía de implementación de la política. Evidencia: pág. 6; E-10.'),
'1.04.c.':('NO CUMPLE','La política contiene objetivo, alcance, entrada en vigor y monitoreo, pero no declara nivel de conformidad y no se identifica un preámbulo; el conjunto obligatorio está incompleto. Evidencia: págs. 3 y 6; E-10.'),
'1.04.d.':('CUMPLE','POL.DGCP.DC.002 §3.6 inventaría portal institucional, Transparencia y vista pública de Comunidad Compras Dominicana. Evidencia: pág. 5; E-10.'),
'1.04.e.':('CUMPLE','POL.DGCP.DC.002 §3.12 establece vigencia anual y revisión/actualización al cumplirse la fecha. Evidencia: pág. 6; E-10.'),
'1.04.f.':('PENDIENTE DE VERIFICACIÓN','Se menciona contratación de alojamiento conforme al SNCCP (§3.4), pero no se aportaron pliegos/contratos de desarrollo externo que exijan accesibilidad; el supuesto y la cláusula deben verificarse. E-10.'),
'2.01.':('NO CUMPLE','La muestra contiene infracciones Nivel A reproducibles (alternativas textuales, teclado, relaciones, propósito de enlaces, nombre/función/valor); por tanto no alcanza Nivel A y menos los niveles superiores. E-01 a E-09.'),
'2.02.':('NO CUMPLE','Las infracciones se encuentran en páginas completas y componentes compartidos de las 11 rutas; no puede afirmarse conformidad de todas las páginas/secciones. E-01 a E-09.'),
'2.03.':('PENDIENTE DE VERIFICACIÓN','No se identificó dentro del dominio un proceso público multietapa completo comparable al ejemplo normativo; los trámites enlazados conducen a sistemas externos excluidos. Requiere inventario institucional de procesos.'),
'2.04.':('NO CUMPLE','Se usan PDFs sin etiquetado/StructTreeRoot y componentes con nombres programáticos ausentes. La muestra documental no es accesiblemente soportada. E-06, E-11.'),
'2.05.':('PENDIENTE DE VERIFICACIÓN','No se ejecutó una matriz completa con JavaScript, traductor, chat y widget de accesibilidad habilitados/inhabilitados/no soportados. Axe sí detectó iframes y controles de terceros sin nombre; se requiere prueba de no interferencia.'),
'3.01.1.a.':('NO CUMPLE','Imágenes funcionales y de contenido carecen de alt (p. ej. servicios_portal.jpg, Fachada-DGCP.jpg, banderas y plugin-pdf.png); existen enlaces gráficos sin nombre. E-01, E-02, E-09.'),
'3.01.2.a.':('CUMPLE','En la muestra, el audiolibro de la Ley 47-25 tiene controles nativos, no usa autoplay y ofrece inmediatamente la ley en texto descargable. El PDF no está etiquetado, deficiencia evaluada en 2.04/3.01.3.a. E-05, E-11.'),
'3.01.2.b.':('PENDIENTE DE VERIFICACIÓN','No se localizó contenido multimedia sincronizado institucional representativo en las 11 páginas que permitiera comprobar subtítulos. No se clasifica como NO APLICA sin inventario completo.'),
'3.01.2.c.':('PENDIENTE DE VERIFICACIÓN','No se localizó video pregrabado institucional representativo para comprobar audiodescripción o alternativa. No se clasifica como NO APLICA sin inventario completo.'),
'3.01.3.a.':('NO CUMPLE','Hay jerarquías de encabezados saltadas, listas inválidas, controles sin etiquetas y cuatro PDFs sin árbol de estructura. E-01, E-02, E-06, E-11.'),
'3.01.3.b.':('PENDIENTE DE VERIFICACIÓN','Se registró el orden secuencial de foco, pero no se verificó manualmente la secuencia de lectura completa con lector de pantalla ni la correspondencia visual/DOM en todas las plantillas.'),
'3.01.3.c.':('PENDIENTE DE VERIFICACIÓN','No se encontró una infracción concreta en la muestra, pero no se revisó exhaustivamente toda instrucción del dominio. No se infiere cumplimiento global por ausencia en la muestra.'),
'3.01.4.a.':('NO CUMPLE','axe/Lighthouse identificaron enlaces dentro de texto distinguibles solo por color y contrastes que agravan esa dependencia en inicio, contacto y publicaciones. E-01 y E-02.'),
'3.01.4.b.':('CUMPLE','El audio institucional localizado no se reproduce automáticamente y expone controles nativos; no se detectó audio automático >3 s en las 11 ejecuciones axe. E-01, E-05.'),
'3.02.1.a.':('NO CUMPLE','Controles declarados role="button" no responden a Espacio: “Nosotros” no expande (aria-expanded permanece false) y #btn1 no cambia de estado; Enter en “Nosotros” navega en vez de abrir el submenú. E-04.'),
'3.02.1.b.':('PENDIENTE DE VERIFICACIÓN','Tab/Shift+Tab no mostraron trampa en la cabecera y se recorrieron hasta 45 paradas por página, pero no se activaron todos los widgets de terceros ni se completó cada página hasta volver al inicio. E-03, E-04.'),
'3.02.2.a.':('PENDIENTE DE VERIFICACIÓN','No se observaron límites de tiempo en la muestra; falta inventario funcional completo para demostrar que no existen en el alcance.'),
'3.02.2.b.':('CUMPLE','La portada ofrece control programático “Pausar animaciones” y controles anterior/siguiente para el carrusel. Es activable con Enter; la falla con Espacio se registra bajo 3.02.1.a. E-04, E-05.'),
'3.02.3.a.':('PENDIENTE DE VERIFICACIÓN','No se ejecutó análisis fotométrico/frecuencia cuadro a cuadro. No se observó destello evidente en capturas, lo cual no basta para dictamen.'),
'3.02.4.a.':('CUMPLE','Existe <a class="skip-link" href="#informacion">Ir al contenido principal</a>; el destino <main id="informacion"> existe y el enlace fue la primera parada en la ejecución limpia. E-04, E-05.'),
'3.02.4.b.':('CUMPLE','Las 11 rutas devuelven títulos no vacíos y descriptivos de la página/consulta. E-03 y E-07.'),
'3.02.4.c.':('PENDIENTE DE VERIFICACIÓN','Tab y Shift+Tab conservaron el orden en la cabecera; no se recorrió manualmente la totalidad de cada DOM ni todos los estados dinámicos. E-03, E-04.'),
'3.02.4.d.':('NO CUMPLE','Hay enlaces sin nombre discernible y 24 enlaces “Más información” en la portada con nombres accesibles iguales y destinos distintos; varios carecen de contexto programático suficiente. E-01, E-05, E-09.'),
'3.03.1.a.':('CUMPLE','Las 11 páginas declaran idioma principal válido: lang="es" o lang="es-es". E-03.'),
'3.03.2.a.':('PENDIENTE DE VERIFICACIÓN','El recorrido de foco no produjo cambio de contexto en la cabecera, pero no se probaron todos los controles/widgets del dominio en todos sus estados. E-03, E-04.'),
'3.03.2.b.':('PENDIENTE DE VERIFICACIÓN','El buscador cambia de contexto tras Enter, comportamiento esperado de envío; no se probaron todas las entradas/selecciones ni sus avisos previos. E-04.'),
'3.03.3.a.':('CUMPLE','El formulario de contacto marca nombre/correo con aria-invalid=true, enlaza mensajes mediante aria-describedby y muestra texto específico para vacío/correo inválido. E-04 y captura contact-errors.png.'),
'3.03.3.b.':('NO CUMPLE','Aunque contacto tiene <label>, el buscador usa solo placeholder y los select de Noticias/Publicaciones carecen de label, aria-label y aria-labelledby. E-01, E-02, E-09.'),
'3.04.1.a.':('NO CUMPLE','El DOM renderizado contiene IDs duplicados compartidos (collapseMain, collapsePortales, collapseInternos, current-lang) y decenas adicionales en Transparencia; la lista del carrusel contiene button como hijos directos de ol. E-03, E-09.'),
'3.04.1.b.':('NO CUMPLE','Control de menú sin nombre, selects sin nombre, enlaces gráficos vacíos e iframe de Google Maps sin title impiden determinar nombre/función programáticamente. E-01, E-02, E-09.'),
}

URLS=[
('Inicio','https://www.dgcp.gob.do/','Portada, carrusel, audiolibro, servicios, noticias, encabezado/pie'),
('Sobre nosotros','https://www.dgcp.gob.do/sobre-nosotros/','Plantilla institucional'),
('Marco legal','https://www.dgcp.gob.do/sobre-nosotros/marco-legal/','Tarjetas y enlaces “Más Información”'),
('Servicios','https://www.dgcp.gob.do/servicios/','Catálogo de servicios'),
('Publicaciones','https://www.dgcp.gob.do/publicaciones/','Buscador/filtros y documentos'),
('Noticias','https://www.dgcp.gob.do/noticias/','Listado y filtros'),
('Contacto','https://www.dgcp.gob.do/contacto/','Formulario, validación y mapa'),
('Transparencia','https://www.dgcp.gob.do/transparencia/','Portada y navegación lateral'),
('Publicaciones oficiales','https://www.dgcp.gob.do/transparencia/publicaciones-oficiales','Acordeones y PDFs'),
('Formularios de servicios','https://www.dgcp.gob.do/sobre-nosotros/marco-legal/formularios-estandar-de-servicios/','Documentos/formularios'),
('Buscador','https://www.dgcp.gob.do/?s=contrataciones','Resultados de búsqueda'),
]

FINDINGS=[
('F-01','CRÍTICA','3.01.1.a.; 3.02.4.d.; 3.04.1.b.','Imágenes y enlaces gráficos sin alternativa/nombre','Inicio; Contacto; Transparencia; todas las plantillas (banderas)','<img src=".../servicios_portal.jpg" class="img-fluid">; <img src=".../Fachada-DGCP.jpg" loading="lazy">','Las imágenes funcionales y de contenido no tienen alt; el enlace que envuelve servicios_portal.jpg queda sin nombre.','Lectores de pantalla no anuncian propósito o contenido; el enlace resulta inutilizable.','axe-core image-alt/link-name; Lighthouse; inspección DOM','Abrir URL, ejecutar axe; localizar los selectores registrados en axe/*.json.','Añadir alt equivalente a imágenes informativas/funcionales; alt="" y role="presentation" solo si son decorativas; nombrar el enlace por texto/aria-label.','E-01, E-02, E-09'),
('F-02','CRÍTICA','3.04.1.b.','Control compartido sin nombre accesible','Las 11 URL','<a class="menu-item dropdown-toggle p-0" ... role="button" aria-controls="collapseMain">','El disparador del menú de cuadrícula no tiene texto, aria-label, aria-labelledby ni title.','Usuario de lector de pantalla encuentra “botón” sin propósito.','axe-core aria-command-name en 11/11; Lighthouse 6/6','Ejecutar axe y filtrar aria-command-name.','Usar <button type="button" aria-label="Abrir enlaces de interés" ...> y mantener aria-expanded sincronizado.','E-01, E-02'),
('F-03','ALTA','3.03.3.b.; 3.04.1.b.','Campos y filtros sin etiqueta programática','Inicio, Noticias, Publicaciones, Contacto','<select id="categoria" ...>; <select name="tags_publicaciones" ...>; <input name="s" placeholder="¿Qué deseas buscar?">','Tres select de Noticias, dos de Publicaciones y buscadores dependen de placeholder/opción visible.','No se comunica propósito fiable al lector de pantalla; dictado por voz no dispone de etiqueta estable.','axe-core select-name; Lighthouse; DOM','Ejecutar axe en /noticias/ y /publicaciones/; inspeccionar labels/ARIA.','Agregar <label for> visible; si el diseño lo exige, etiqueta visualmente oculta, no solo placeholder.','E-01, E-02, E-09'),
('F-04','ALTA','3.02.1.a.','Activación incompleta por teclado en controles con role=button','Inicio/encabezado compartido','<a ... role="button">Nosotros...</a>; <a id="btn1" ... type="button">','Espacio no cambia aria-expanded del menú ni el estado/título de #btn1; Enter en “Nosotros” navega a otra URL.','Personas que usan teclado no pueden operar la función expuesta como botón de forma equivalente.','Playwright: Enter/Espacio, lectura de URL, aria-expanded y title','Enfocar “Nosotros”, pulsar Espacio: aria-expanded sigue false. Enfocar #btn1 y pulsar Espacio: title no cambia.','Usar <button>; manejar Enter y Espacio; separar enlace de navegación del botón que expande; probar estados y aria-expanded.','E-04'),
('F-05','ALTA','3.02.4.d.','Enlaces ambiguos “Más información” y enlaces vacíos','Inicio y Marco legal','<a ...>Más información</a> (24 en inicio); <a href="...Login.aspx"><img ... sin alt></a>','Muchos enlaces comparten nombre accesible genérico con destinos diferentes; otros no tienen nombre.','La lista de enlaces de un lector de pantalla no permite escoger destino.','DOM renderizado; axe link-name; Lighthouse','Listar enlaces cuyo texto coincide con “Más información”; comparar href y nombre accesible.','Incluir destino en el nombre visible: “Más información sobre …”; asociar mediante aria-labelledby al título de la tarjeta.','E-01, E-05, E-09'),
('F-06','ALTA','3.01.3.a.; 3.04.1.a.','Estructura HTML inconsistente e IDs duplicados','Todas; especialmente Inicio y Transparencia','<ol class="carousel-indicators"><button ...>; ids collapseMain/current-lang repetidos','Saltos H1-H4/H5, lista con hijos button y IDs repetidos comprometen relaciones y parseo.','La navegación por encabezados y las relaciones ARIA/fragmentos pueden ser ambiguas.','axe list/heading-order; inspección DOM duplicateIds','Ejecutar recolector y revisar dom/*.json: headings y duplicateIds.','Jerarquía H1→H2→H3; envolver botones en li o usar div; generar IDs únicos y ajustar referencias.','E-01, E-02, E-03, E-09'),
('F-07','ALTA','2.04.; 3.01.3.a.','PDF públicos sin etiquetado estructural','Ley 47-25; revista 2025; estudio mujeres MIPYMES; formulario RPE','/MarkInfo ausente o /Marked=false; /StructTreeRoot ausente en 4/4','Ninguno de los cuatro documentos tiene etiquetas; tres no declaran /Lang y la revista declara en-US.','Lectura lineal, encabezados, listas, tablas y orden no se transmiten correctamente a tecnología asistiva.','pypdf 6.1.1 y pdfinfo; inspection.json','Descargar archivos listados y ejecutar inspect-pdfs.py/pdfinfo -meta.','Remediar en archivo fuente/PDF: etiquetas, orden, encabezados, alt, idioma es-DO, título y campos etiquetados; validar con PAC/Acrobat.','E-11'),
('F-08','ALTA','3.01.4.a.','Enlaces diferenciados solo por color','Inicio, Contacto, Publicaciones','Enlaces dentro de bloques de texto sin subrayado/indicador no cromático','axe reporta link-in-text-block; Lighthouse registra contrastes deficientes, incluido 2.67:1 en breadcrumb.','Personas con deficiencia de percepción cromática pueden no reconocer enlaces.','axe-core link-in-text-block; Lighthouse color-contrast','Abrir reporte Lighthouse/axe y localizar selectores.','Subrayar enlaces o añadir otro indicador persistente; corregir colores. El contraste es además deuda AA aunque no sea el único fundamento A.','E-01, E-02'),
('F-09','MEDIA','3.01.1.a.; 3.04.1.b.','Iframe de mapa sin nombre','Contacto','<iframe class="embed-responsive-item" src="https://www.google.com/maps/embed?..." allowfullscreen>','El iframe no tiene title/aria-label.','El lector de pantalla no puede identificar el contenido incrustado.','axe frame-title; Lighthouse','Ejecutar axe en /contacto/.','Añadir title="Mapa de ubicación de la DGCP en Santo Domingo".','E-01, E-02'),
('F-10','MEDIA','1.04.; 1.04.c.; 1.04.f.','Política interna incompleta frente al formulario','Documento POL.DGCP.DC.002','§3.8–3.13; sin nivel de conformidad declarado ni cláusula contractual aportada','El documento prueba política/responsable/inventario/revisión, pero no todos los elementos obligatorios.','Riesgo de gobernanza y adquisición de componentes inaccesibles.','Revisión visual de las 8 páginas; documento escaneado sin texto extraíble','Consultar páginas 3, 5, 6 y 8 de la política adjunta.','Actualizar política con preámbulo, nivel objetivo y criterios contractuales verificables; anexar pliegos/contratos.','E-10'),
]

EVIDENCE=[
('E-01','axe-core JSON','audit/evidence/axe/*.json','11 páginas; reglas, impacto, selectores y HTML de cada nodo.'),
('E-02','Lighthouse JSON/HTML','audit/evidence/lighthouse/*.{json,html}','6 páginas; puntuaciones 0.76–0.87 y auditorías fallidas.'),
('E-03','DOM y teclado','audit/evidence/dom/*.json','DOM renderizado, headings, imágenes, enlaces, controles, tablas, idiomas, IDs y hasta 45 paradas Tab.'),
('E-04','Pruebas funcionales','audit/evidence/functional-*.json; audit/evidence/keyboard-detail.json','Tab/Shift+Tab, Enter/Espacio, menú, buscador, formulario y errores.'),
('E-05','HTML renderizado','audit/evidence/html/*.html','Fragmentos exactos; audiolibro, skip link, carrusel, enlaces.'),
('E-06','Capturas','audit/evidence/screenshots/*.png','Capturas de las 11 páginas y errores de contacto.'),
('E-07','Inventario','Hoja Inventario URL del Excel','URLs, plantilla y resultados básicos.'),
('E-08','Código reproducible','audit/collect-evidence.mjs; run-lighthouse.mjs; functional-tests.mjs; keyboard-detail.mjs','Scripts utilizados.'),
('E-09','Consolidado técnico','audit/evidence/consolidated-findings.json','Hallazgos y nodos representativos.'),
('E-10','Política/procedimiento internos','POL.DGCP.DC.002…pdf; PRO.DGCP.DDIS.001…pdf','Evidencia documental para serie 1.04.'),
('E-11','PDF públicos','audit/evidence/pdfs/inspection.json y 4 PDF','Etiquetado, idioma, metadatos, texto y campos.'),
]

def criterion_rows(wb):
    for sn in ['Políticas Internas','Dir. de Conformidad','Ppio. Perceptible','Ppio. Operable','Ppio. Comprensible','Ppio. Robusto']:
        ws=wb[sn]
        for r in range(1,ws.max_row+1):
            for c in range(2,5):
                v=ws.cell(r,c).value
                if isinstance(v,str) and v in S:
                    yield ws,r,c,v

wb=load_workbook(SOURCE)
found=set()
for ws,r,c,cid in criterion_rows(wb):
    state,comment=S[cid]; found.add(cid)
    state_col=5 if ws.title in ['Políticas Internas','Dir. de Conformidad'] else 7
    comment_col=state_col+1
    ws.cell(r,state_col,state); ws.cell(r,comment_col,comment)
    ws.cell(r,state_col).alignment=Alignment(wrap_text=True,vertical='top')
    ws.cell(r,comment_col).alignment=Alignment(wrap_text=True,vertical='top')
assert found==set(S), (set(S)-found, found-set(S))

# Convierte las hojas preliminares heredadas en resúmenes finales, manteniendo sus nombres y estructura.
pre=wb['Evaluación preliminar DGCP']; pre['D1']='Dictamen final'; pre['E1']='Justificación final / evidencia'
for r in range(2,pre.max_row+1):
    cid=pre.cell(r,2).value
    if cid in S:
        pre.cell(r,4,S[cid][0]); pre.cell(r,5,S[cid][1]); pre.cell(r,6,'Véase Índice de evidencia y Hallazgos técnicos')
scope=wb['Alcance y metodología DGCP']
scope['A1']='EVALUACIÓN TÉCNICA FINAL · NORTIC B2:2017 · NIVEL A'
scope['B7']='Auditoría técnica por muestreo representativo con DOM renderizado, axe-core, Lighthouse, teclado, formularios y PDF.'
scope['B8']='CUMPLE solo con evidencia suficiente; NO CUMPLE ante infracción reproducible; PENDIENTE cuando falta evidencia; NO APLICA solo con justificación de alcance.'
scope['B15']='Remediar hallazgos y completar lector de pantalla real, WAVE/PAC, destellos, procesos y pruebas sin tecnologías no esenciales.'
res=wb['Resultado auditoría 14-10-2026']
summary_rows={
  1:('EVALUACIÓN TÉCNICA FINAL DGCP.GOB.DO — NORTIC B2:2017 A',None),
  2:('Alcance','11 URL/plantillas de www.dgcp.gob.do; 4 PDF; políticas internas. SECP y sistemas externos excluidos.'),
  3:('Fecha','2026-10-14'),
  4:('Método','DOM renderizado, axe-core 4.14.0, Lighthouse 12.8.2, Chromium 149, Playwright 1.64.1, teclado y pypdf 6.1.1.'),
  5:('CUMPLE','11'), 6:('NO CUMPLE','13'), 7:('NO APLICA','0'), 8:('PENDIENTE DE VERIFICACIÓN','13'),
  9:('Dictamen','NO CONFORME en la muestra: existen infracciones Nivel A reproducibles. No equivale a certificación OGTIC.'),
  10:('Criterios','37 identificadores exactos del formulario.'),
  11:('Evidencia automatizada','audit/evidence/axe y audit/evidence/lighthouse'),
  12:('Evidencia DOM/funcional','audit/evidence/dom, functional-*.json y keyboard-detail.json'),
  13:('Evidencia PDF','audit/evidence/pdfs/inspection.json'),
  14:('Informe','Informe_Auditoria_Tecnica_DGCP_NORTIC_B2_A_2026-10-14.md'),
  15:('Limitaciones','Sin WAVE, lector de pantalla real, PAC/Acrobat ni análisis fotométrico de destellos.'),
}
for r,(a,b) in summary_rows.items(): res.cell(r,1,a); res.cell(r,2,b)

for name in ['Matriz final DGCP','Inventario URL','Hallazgos técnicos','Índice de evidencia','Herramientas y límites']:
    if name in wb.sheetnames: del wb[name]

colors={'CUMPLE':'C6EFCE','NO CUMPLE':'FFC7CE','NO APLICA':'D9EAF7','PENDIENTE DE VERIFICACIÓN':'FFEB9C'}
thin=Side(style='thin',color='B7C9DC')
header=PatternFill('solid',fgColor='003876')
def style_sheet(ws,widths):
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for c in ws[1]: c.fill=header;c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
    for row in ws.iter_rows():
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top');c.border=Border(bottom=thin)
    for col,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(col)].width=w

matrix=wb.create_sheet('Matriz final DGCP')
matrix.append(['Área','Identificador','Requisito normativo (sin modificar)','Estado','Justificación y evidencia'])
for ws,r,c,cid in criterion_rows(wb):
    req_col=3 if ws.title=='Políticas Internas' else 4
    requirement=ws.cell(r,req_col).value
    if ws.title.startswith('Ppio.'):
        crit=ws.cell(r,4).value; directive=ws.cell(r,5).value
        requirement=f'{crit}: {directive}'
    state,comment=S[cid]; matrix.append([ws.title,cid,requirement,state,comment]); matrix.cell(matrix.max_row,4).fill=PatternFill('solid',fgColor=colors[state])
style_sheet(matrix,[23,16,85,30,105])

inv=wb.create_sheet('Inventario URL'); inv.append(['Plantilla','URL','Cobertura','HTTP','lang','Título','axe: reglas','axe: nodos','Captura'])
for name,url,cov in URLS:
    key={'Sobre nosotros':'sobre-nosotros','Marco legal':'marco-legal','Publicaciones oficiales':'transparencia-publicaciones','Formularios de servicios':'formularios-servicios','Buscador':'buscar'}.get(name,name.lower())
    dom=json.loads((EV/'dom'/f'{key}.json').read_text())['dom']; axe=json.loads((EV/'axe'/f'{key}.json').read_text())
    inv.append([name,url,cov,200,dom.get('htmlLang'),dom.get('title'),len(axe['violations']),sum(len(x['nodes']) for x in axe['violations']),f'audit/evidence/screenshots/{key}.png'])
style_sheet(inv,[28,65,60,10,12,70,14,14,50])

fh=wb.create_sheet('Hallazgos técnicos'); fh.append(['ID','Prioridad','Criterio NORTIC','Hallazgo','URL/alcance','Elemento/fragmento HTML','Incumplimiento','Impacto','Herramienta/procedimiento','Reproducción','Solución recomendada','Evidencia'])
for row in FINDINGS: fh.append(row)
style_sheet(fh,[10,14,24,35,45,60,65,55,40,60,65,20])

ei=wb.create_sheet('Índice de evidencia');ei.append(['ID','Tipo','Ruta','Descripción'])
for row in EVIDENCE:ei.append(row)
style_sheet(ei,[10,28,70,90])

tl=wb.create_sheet('Herramientas y límites');tl.append(['Campo','Detalle'])
tools=[
('Fecha y zona','2026-10-14 · America/Santo_Domingo'),('Chromium','149.0.7827.102'),('Playwright','1.64.1'),('axe-core','4.14.0'),('Lighthouse','12.8.2'),('pypdf','6.1.1'),('Muestra web','11 URLs/plantillas; DOM renderizado con JavaScript.'),('Lighthouse','6 URLs; accesibilidad escritorio.'),('PDF','4 archivos públicos representativos.'),('Teclado','Tab, Shift+Tab, Enter y Espacio; hasta 45 paradas por URL y pruebas dirigidas.'),('No ejecutado','WAVE (no disponible); lector de pantalla real NVDA/JAWS/VoiceOver; PAC/Acrobat; análisis fotométrico de destellos; contraste manual de todos los estados.'),('Límite de conformidad','El muestreo no demuestra conformidad integral del dominio. Los criterios de página/proceso completo se califican NO CUMPLE cuando existe infracción y PENDIENTE cuando falta evidencia.'),('Sistemas excluidos','SECP y otros sistemas/subdominios externos, aunque estén enlazados.'),]
for row in tools:tl.append(row)
style_sheet(tl,[28,120])

# Actualiza catálogo de estados sin alterar criterios.
est=wb['Estados']; est['C3']='CUMPLE'; est['C4']='NO CUMPLE'; est['C5']='NO APLICA'; est['C6']='PENDIENTE DE VERIFICACIÓN'
wb.save(OUTPUT)

# Consolidado JSON reproducible.
consolidated={'date':'2026-10-14','statusCounts':dict(Counter(x[0] for x in S.values())),'criteria':{k:{'status':v[0],'justification':v[1]} for k,v in S.items()},'findings':[dict(zip(['id','priority','criteria','title','scope','html','description','impact','method','reproduce','solution','evidence'],x)) for x in FINDINGS]}
(EV/'consolidated-findings.json').write_text(json.dumps(consolidated,ensure_ascii=False,indent=2),encoding='utf-8')

axe_rows=[]
for p in sorted((EV/'axe').glob('*.json')):
    j=json.loads(p.read_text());axe_rows.append((p.stem,len(j['violations']),sum(len(v['nodes']) for v in j['violations']),', '.join(v['id'] for v in j['violations'])))
lh_rows=[]
for p in sorted((EV/'lighthouse').glob('*.json')):
    j=json.loads(p.read_text()); failed=[a['id'] for a in j['categories']['accessibility']['auditRefs'] if j['audits'][a['id']].get('scoreDisplayMode')=='binary' and j['audits'][a['id']].get('score')==0]
    lh_rows.append((p.stem,j['categories']['accessibility']['score'],', '.join(failed)))
counts=Counter(x[0] for x in S.values())

report=f'''# Auditoría técnica de accesibilidad web — DGCP / NORTIC B2:2017 Nivel A

**Fecha de prueba:** 14 de octubre de 2026 (America/Santo_Domingo)  
**Dominio incluido:** `https://www.dgcp.gob.do/` y rutas públicas del mismo portal  
**Excluido:** SECP y sistemas/subdominios externos independientes.

## Resumen ejecutivo

El portal **no demuestra conformidad NORTIC B2:2017 Nivel A** en el alcance auditado. De los 37 identificadores exactos del formulario: **{counts['CUMPLE']} CUMPLE**, **{counts['NO CUMPLE']} NO CUMPLE**, **{counts['PENDIENTE DE VERIFICACIÓN']} PENDIENTE DE VERIFICACIÓN** y **0 NO APLICA**. La decisión no se basa en puntuaciones: existen infracciones Nivel A concretas y reproducibles en componentes compartidos y contenido público.

Riesgos principales: imágenes/enlaces gráficos sin alternativa; controles, filtros y enlaces sin nombre; operación incompleta por teclado; enlaces “Más información” ambiguos; estructura/IDs defectuosos; enlaces distinguibles solo por color; y PDF públicos sin etiquetas. La portada obtuvo 0.76 en Lighthouse Accessibility y fue la plantilla con mayor concentración de incidencias. Las demás páginas Lighthouse quedaron entre 0.81 y 0.87; esas cifras **no equivalen a conformidad**.

Aspectos verificados positivamente: las 11 páginas declaran idioma principal; sus títulos son descriptivos; existe enlace de salto funcional; el formulario de contacto asocia etiquetas y errores de texto; el audiolibro usa controles nativos, no autoplay y ofrece la ley en texto; la política interna aporta responsable, inventario y revisión anual.

## Alcance real e inventario

| Plantilla | URL | Cobertura |
|---|---|---|
'''+''.join(f'| {n} | {u} | {c} |\n' for n,u,c in URLS)+f'''

Además se inspeccionaron cuatro PDF representativos: Ley 47-25 (166 páginas), Revista Contrataciones 2025 (76), estudio Mujeres MIPYMES (76) y formulario RPE (2). Se revisaron los documentos internos `POL.DGCP.DC.002` y `PRO.DGCP.DDIS.001` para la serie 1.04.

## Metodología y herramientas efectivamente utilizadas

- Chromium 149.0.7827.102 + Playwright 1.64.1: DOM renderizado, estados, formularios, teclado y capturas.
- axe-core 4.14.0: 11 páginas, reglas WCAG A y buenas prácticas; JSON completo exportado.
- Lighthouse 12.8.2: 6 plantillas, JSON y HTML exportados.
- Inspección DOM propia: idiomas, títulos, headings, landmarks, imágenes, enlaces, controles, ARIA, tablas, medios, contenido oculto e IDs.
- pypdf 6.1.1 y `pdfinfo`: estructura, etiquetas, idioma, metadatos, formularios y texto extraíble de PDF.
- Revisión visual de las ocho páginas escaneadas de la política interna.
- Contraste: auditorías automatizadas axe/Lighthouse; no se completó medición manual de todos los estados.

No se ejecutó WAVE porque no estaba disponible. Tampoco se usó lector de pantalla real, PAC/Acrobat, análisis cuadro a cuadro de destellos ni una prueba humana completa de orden de lectura. Estos límites se reflejan como **PENDIENTE DE VERIFICACIÓN**, no como cumplimiento ni NO APLICA.

## Resultados automatizados

### axe-core

| Página | Reglas con infracción | Nodos | IDs de regla |
|---|---:|---:|---|
'''+''.join(f'| {n} | {r} | {nodes} | {ids} |\n' for n,r,nodes,ids in axe_rows)+'''

### Lighthouse Accessibility

| Página | Puntuación | Auditorías binarias fallidas |
|---|---:|---|
'''+''.join(f'| {n} | {score:.2f} | {failed} |\n' for n,score,failed in lh_rows)+'''

## Pruebas funcionales

- `Tab` y `Shift+Tab`: la ejecución limpia situó primero “Ir al contenido principal”; ida y vuelta conservaron la secuencia de cabecera. En el muestreo ampliado se registraron hasta 45 paradas por URL sin trampa observada, pero no se activaron todos los widgets.
- Menú “Nosotros”: Espacio no lo expande (`aria-expanded=false`); Enter navega a `/sobre-nosotros/` en vez de abrir el submenú. Incumple 3.02.1.a para un elemento declarado `role="button"`.
- Control `#btn1` “Pausar animaciones”: Enter cambia su estado/título; Espacio no lo hace. Incumple 3.02.1.a.
- Buscador: al introducir `prueba-accesibilidad-zzzx` y pulsar Enter, carga la URL de resultados esperada. La entrada, sin embargo, carece de etiqueta programática estable.
- Contacto: con nombre vacío y correo inválido aparecen mensajes específicos; los campos reciben `aria-invalid=true` y `aria-describedby`. El foco permanece en el envío, pero 3.03.3.a exige identificación y descripción, que sí se verificaron.
- Audiolibro Ley 47-25: `<audio controls>` con MP3/OGG, sin autoplay, y enlace adyacente a la ley. El PDF alternativo no está etiquetado.

## Matriz de cumplimiento

| ID | Estado | Justificación resumida |
|---|---|---|
'''+''.join(f'| {cid} | {state} | {comment} |\n' for cid,(state,comment) in S.items())+'''

## Incumplimientos reproducibles

'''+''.join(f'''### {fid} — {title} ({prio})

**Criterio:** {criteria}  
**URL/alcance:** {scope}  
**Elemento:** `{html}`  
**Incumplimiento:** {desc}  
**Impacto:** {impact}  
**Procedimiento:** {method}  
**Reproducción:** {repro}  
**Solución:** {solution}  
**Evidencia:** {ev}

''' for fid,prio,criteria,title,scope,html,desc,impact,method,repro,solution,ev in FINDINGS)+'''
## Criterios pendientes

Permanecen pendientes 1.04.f, 2.03, 2.05, 3.01.2.b, 3.01.2.c, 3.01.3.b, 3.01.3.c, 3.02.1.b, 3.02.2.a, 3.02.3.a, 3.02.4.c, 3.03.2.a y 3.03.2.b. Requieren, según el caso: contratos/pliegos, inventario de procesos, matriz sin JavaScript/terceros, inventario audiovisual, lector de pantalla real, recorrido completo de foco y análisis de destellos.

## Plan de remediación priorizado

1. **0–15 días — bloqueo Nivel A:** corregir nombres/alt, etiquetas de filtros/buscadores, enlaces vacíos, botón de menú compartido, soporte Enter/Espacio y enlaces “Más información”.
2. **0–30 días — estructura:** IDs únicos, headings secuenciales, lista del carrusel válida, iframe con título y enlaces no dependientes del color.
3. **15–60 días — documentos:** remediar primero Ley 47-25, formularios y publicaciones más descargadas; exigir PDF/UA o estructura equivalente en el flujo editorial.
4. **30–60 días — gobernanza:** actualizar POL.DGCP.DC.002 con nivel de conformidad y cláusulas de compras/desarrollo accesible; conservar evidencias de revisión.
5. **Tras correcciones:** repetir los scripts, prueba humana completa con NVDA/Firefox y VoiceOver/Safari, PAC/Acrobat, contraste manual y pruebas sin tecnologías no esenciales.

## Trazabilidad de evidencia

'''+''.join(f'- **{i} — {t}:** `{p}` — {d}\n' for i,t,p,d in EVIDENCE)+'''

## Limitación del dictamen

Este informe es una auditoría técnica por muestreo representativo, no una certificación OGTIC ni una afirmación de cobertura exhaustiva de cada URL histórica. Un solo fallo comprobado basta para **NO CUMPLE**; la ausencia de fallo en la muestra no se convirtió automáticamente en CUMPLE. Los criterios de alcance/proceso o prueba humana insuficiente se mantienen pendientes.
'''
REPORT.write_text(report,encoding='utf-8')
print(json.dumps({'excel':str(OUTPUT),'report':str(REPORT),'counts':dict(counts),'criteria':len(S),'findings':len(FINDINGS)},ensure_ascii=False,indent=2))
